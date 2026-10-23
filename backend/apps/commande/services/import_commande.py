"""Import d'une commande depuis un fichier Excel.

Modèle (un seul onglet) :

    Objet     | <objet de la demande>
    (vide)
    Article   | Quantité | Bénéficiaire
    Stylo     | 1        | RAKOTO Jean
              | 2        | DSI            <- article vide : même article que la ligne au-dessus
    Ramette   | 5        |                <- bénéficiaire vide : pour le demandeur

- Article : code ou désignation d'un article du catalogue ; sinon, demande de **nouvel article** (le texte
  devient la spécification). Vide : l'article de la ligne précédente (cellules fusionnées comprises).
- Bénéficiaire : vide = le demandeur ; sinon matricule ou nom d'un employé (de la direction du demandeur),
  la direction du demandeur, un site ou une salle.
- Les lignes entièrement vides sont ignorées ; un même article (sur plusieurs lignes ou plusieurs blocs) forme
  une seule ligne de commande, répartie entre ses bénéficiaires (quantités additionnées par bénéficiaire).
"""
import io
from collections import OrderedDict
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.catalogue.models import Article
from apps.common.imports import normaliser_texte
from apps.employee.models import Direction, Employer, Site
from apps.stock.models import Salle

ENTETES = ("Article", "Quantité", "Bénéficiaire")
LIGNES_MODELE = 100
LIGNES_MAX = 1000
OBJET_MAX = 100
SPECIFICATION_MAX = 255


# --------------------------------------------------------------------------
# Modèle à télécharger
# --------------------------------------------------------------------------
def generer_modele():
    """Octets du modèle vide : « Objet » en ligne 1, en-tête « Article | Quantité | Bénéficiaire » en ligne 3."""
    bleu, jaune = "0B4F8A", "FFF4CC"
    fin = Side(style="thin", color="C9D2DE")
    bord = Border(left=fin, right=fin, top=fin, bottom=fin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Commande"
    ws["A1"] = "Objet"
    ws["A1"].font = Font(name="Arial", bold=True)
    ws.merge_cells("B1:C1")
    for col in "BC":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=jaune)
        ws[f"{col}1"].border = bord
        ws[f"{col}1"].font = Font(name="Arial")
    for i, titre in enumerate(ENTETES, start=1):
        cellule = ws.cell(row=3, column=i, value=titre)
        cellule.font = Font(name="Arial", bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=bleu)
        cellule.border = bord
    for ligne in range(4, 4 + LIGNES_MODELE):
        for col in (1, 2, 3):
            cellule = ws.cell(row=ligne, column=col)
            cellule.border = bord
            cellule.font = Font(name="Arial")
        ws.cell(row=ligne, column=2).alignment = Alignment(horizontal="center")
    for col, largeur in zip("ABC", (42, 12, 34), strict=True):
        ws.column_dimensions[col].width = largeur
    ws.freeze_panes = "A4"
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


# --------------------------------------------------------------------------
# Lecture du fichier
# --------------------------------------------------------------------------
class FichierCommandeInvalide(ValueError):
    """Fichier qui ne suit pas le modèle (objet ou en-tête introuvable...)."""


def _texte(valeur):
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return " ".join(str(valeur).split())


def lire_fichier(fichier):
    """``(objet, [(n° de ligne Excel, article, quantité, bénéficiaire), ...])`` (textes bruts)."""
    try:
        classeur = load_workbook(fichier, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - fichier illisible, quelle qu'en soit la raison
        raise FichierCommandeInvalide("Impossible de lire le fichier : utilisez le modèle Excel (.xlsx).") from exc
    feuille = classeur.worksheets[0]
    objet, entete, lignes = None, None, []
    for numero, valeurs in enumerate(feuille.iter_rows(max_col=6, values_only=True), start=1):
        textes = [_texte(v) for v in valeurs] + [""] * 3
        if entete is None:
            cles = [normaliser_texte(t) for t in textes[:3]]
            if cles[0] == "OBJET":
                objet = next((t for t in textes[1:] if t), "")
            elif cles == [normaliser_texte(t) for t in ENTETES]:
                entete = numero
            continue
        if len(lignes) >= LIGNES_MAX:
            raise FichierCommandeInvalide(f"Le fichier dépasse {LIGNES_MAX} lignes.")
        lignes.append((numero, textes[0], valeurs[1] if len(valeurs) > 1 else None, textes[2]))
    classeur.close()
    if entete is None:
        raise FichierCommandeInvalide("En-tête « Article | Quantité | Bénéficiaire » introuvable : utilisez le modèle.")
    return objet, lignes


# --------------------------------------------------------------------------
# Correspondances (articles, bénéficiaires)
# --------------------------------------------------------------------------
class Referentiels:
    """Articles et bénéficiaires possibles pour un demandeur, indexés par texte normalisé."""

    def __init__(self, demandeur):
        self.demandeur = demandeur
        self.articles_code = {}
        self.articles_nom = {}
        for article in Article.objects.only("code_article", "designation", "is_immobilisation").order_by("code_article"):
            self.articles_code.setdefault(normaliser_texte(article.code_article), article)
            self.articles_nom.setdefault(normaliser_texte(article.designation), article)

        # Comme dans le formulaire : employés et direction du demandeur, tous les sites et salles.
        direction = getattr(demandeur, "emp_dir_id", None)
        employes = Employer.objects.all()
        if direction is not None:
            employes = employes.filter(emp_dir_id=direction)
        self.employes_matricule, self.employes_nom = {}, {}
        for employe in employes:
            if employe.emp_matricule:
                self.employes_matricule[normaliser_texte(employe.emp_matricule)] = employe
            self.employes_nom.setdefault(normaliser_texte(employe.emp_nom), []).append(employe)
        self.directions = {}
        for d in [direction] if direction is not None else Direction.objects.all():
            for texte in (d.dir_libelle, d.dir_description):
                if texte:
                    self.directions.setdefault(normaliser_texte(texte), d)
        self.sites = {normaliser_texte(s.site_nom): s for s in Site.objects.all()}
        self.salles = {}
        for salle in Salle.objects.select_related("localite"):
            for texte in (salle.nom, f"{salle.nom} ({salle.localite.site_nom})"):
                self.salles.setdefault(normaliser_texte(texte), []).append(salle)

    def article(self, texte):
        cle = normaliser_texte(texte)
        return self.articles_code.get(cle) or self.articles_nom.get(cle)

    def beneficiaire(self, texte):
        """``(type, objet)`` ; lève ``ValueError`` (message pour l'utilisateur) si introuvable ou ambigu."""
        if not texte:
            return "EMPLOYE", self.demandeur
        cle = normaliser_texte(texte)
        # « NOM (matricule) » : le matricule entre parenthèses suffit.
        if texte.endswith(")") and "(" in texte:
            matricule = normaliser_texte(texte[texte.rindex("(") + 1:-1])
            if matricule in self.employes_matricule:
                return "EMPLOYE", self.employes_matricule[matricule]
        if cle in self.employes_matricule:
            return "EMPLOYE", self.employes_matricule[cle]
        homonymes = self.employes_nom.get(cle, [])
        if len(homonymes) > 1:
            raise ValueError(f"Plusieurs employés s'appellent « {texte} » : précisez le matricule.")
        if homonymes:
            return "EMPLOYE", homonymes[0]
        if cle in self.directions:
            return "DIRECTION", self.directions[cle]
        if cle in self.sites:
            return "SITE", self.sites[cle]
        salles = self.salles.get(cle, [])
        if len(salles) > 1:
            raise ValueError(f"Plusieurs salles s'appellent « {texte} » : précisez le site, ex. « {texte} (Siège) ».")
        if salles:
            return "SALLE", salles[0]
        raise ValueError(f"Bénéficiaire « {texte} » introuvable (employé ou direction de votre direction, site ou salle).")


CHAMPS = {"EMPLOYE": "employe_beneficiaire", "DIRECTION": "direction_beneficiaire",
          "SITE": "site_beneficiaire", "SALLE": "salle_beneficiaire"}
LIBELLES_TYPE = {"EMPLOYE": "Employé", "DIRECTION": "Direction", "SITE": "Site", "SALLE": "Salle"}


ATTRIBUTS_NOM = {"EMPLOYE": "emp_nom", "DIRECTION": "dir_libelle", "SITE": "site_nom", "SALLE": "nom"}


def _nom(type_, cible):
    return getattr(cible, ATTRIBUTS_NOM[type_])


def _quantite(valeur):
    try:
        quantite = Decimal(str(valeur).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ValueError("Quantité invalide : nombre entier supérieur à 0 attendu.") from None
    if quantite <= 0 or quantite != quantite.to_integral_value():
        raise ValueError("Quantité invalide : nombre entier supérieur à 0 attendu.")
    return int(quantite)


# --------------------------------------------------------------------------
# Analyse
# --------------------------------------------------------------------------
def analyser(objet, lignes, demandeur):
    """Rapport ligne par ligne et charge utile de la commande (``payload``, ``None`` s'il y a des erreurs)."""
    refs = Referentiels(demandeur)
    rapport, erreurs = [], []
    groupes = OrderedDict()  # article (code ou spécification normalisée) -> ligne de commande
    article_courant = ""

    objet = (objet or "").strip()
    if not objet:
        erreurs.append("Objet de la demande manquant (ligne 1, à côté de « Objet »).")
    elif len(objet) > OBJET_MAX:
        erreurs.append(f"Objet trop long ({len(objet)} caractères, {OBJET_MAX} au plus).")

    for numero, texte_article, valeur_quantite, texte_beneficiaire in lignes:
        if not texte_article and not _texte(valeur_quantite) and not texte_beneficiaire:
            continue  # ligne vide : ignorée (le bloc de l'article courant continue)
        ligne = {"ligne": numero, "article_saisi": texte_article, "article_repris": not texte_article,
                 "quantite": None, "beneficiaire_saisi": texte_beneficiaire}
        rapport.append(ligne)
        try:
            if texte_article:
                article_courant = texte_article
            if not article_courant:
                raise ValueError("Article manquant : indiquez l'article (ou placez cette ligne sous un article).")
            article = refs.article(article_courant)
            if article is None and len(article_courant) > SPECIFICATION_MAX:
                raise ValueError(f"Nouvel article : description trop longue ({SPECIFICATION_MAX} caractères au plus).")
            ligne.update({
                "article": article.code_article if article else None,
                "designation": article.designation if article else article_courant,
                "nouvel_article": article is None,
            })
            ligne["quantite"] = _quantite(valeur_quantite)
            type_, cible = refs.beneficiaire(texte_beneficiaire)
            ligne.update({"beneficiaire_type": LIBELLES_TYPE[type_], "beneficiaire": _nom(type_, cible),
                          "pour_demandeur": not texte_beneficiaire})
        except ValueError as exc:
            ligne.update({"statut": "ERREUR", "message": str(exc)})
            continue
        ligne["statut"] = "OK"
        cle = ("A", article.code_article) if article else ("N", normaliser_texte(article_courant))
        groupe = groupes.setdefault(cle, {"article": article, "specification": article_courant, "attributions": OrderedDict()})
        attribution = groupe["attributions"].setdefault((type_, cible.pk), {"type": type_, "cible": cible, "quantite": 0})
        attribution["quantite"] += ligne["quantite"]

    lignes_ok = sum(1 for ligne in rapport if ligne["statut"] == "OK")
    if not rapport:
        erreurs.append("Aucun article dans le fichier : remplissez au moins une ligne sous l'en-tête.")
    nb_erreurs = len(erreurs) + sum(1 for ligne in rapport if ligne["statut"] == "ERREUR")

    details = []
    for groupe in groupes.values():
        attributions = list(groupe["attributions"].values())
        details.append({
            **({"article": groupe["article"].code_article} if groupe["article"] else {"specification": groupe["specification"]}),
            "quantite": sum(a["quantite"] for a in attributions),
            "attributions": [{CHAMPS[a["type"]]: a["cible"].pk, "quantite": a["quantite"], "quantite_demandee": a["quantite"]}
                             for a in attributions],
        })
    resume = [
        {"article": g["article"].code_article if g["article"] else None,
         "designation": g["article"].designation if g["article"] else g["specification"],
         "nouvel_article": g["article"] is None,
         "quantite": sum(a["quantite"] for a in g["attributions"].values()),
         "repartition": [{"beneficiaire": _nom(a["type"], a["cible"]), "type": LIBELLES_TYPE[a["type"]], "quantite": a["quantite"]}
                         for a in g["attributions"].values()]}
        for g in groupes.values()
    ]
    return {
        "objet": objet,
        "erreurs": erreurs,
        "lignes": rapport,
        "lignes_ok": lignes_ok,
        "lignes_erreur": nb_erreurs - len(erreurs),
        "commande": resume,
        "valide": nb_erreurs == 0,
        "payload": {"objet": objet, "employe_demandeur": demandeur.pk, "details": details} if nb_erreurs == 0 else None,
    }
